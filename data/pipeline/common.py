"""Shared conventions for the per-carrier pipelines.

Column lists mirror docs/data/SCHEMA.md. Every pipeline writes only to
PROCESSED_DIR / <carrier> and reads raw files through `raw_path`, which
refuses files that are not in the manifest or not on disk.
"""
from __future__ import annotations

import csv
import html
import re
from decimal import ROUND_HALF_UP, Decimal
from functools import lru_cache
from pathlib import Path
from typing import Iterable

DATA_DIR = Path(__file__).resolve().parents[1]
RAW_DIR = DATA_DIR / "raw"
PROCESSED_DIR = DATA_DIR / "processed"
MANIFEST = RAW_DIR / "manifest.csv"

CARRIERS = ("usps", "ups", "fedex")
RATE_CARDS = {"usps": ("retail", "commercial"), "ups": ("list",), "fedex": ("list",)}
ZONES = {"usps": range(1, 10), "ups": range(2, 9), "fedex": range(2, 9)}
MAX_WEIGHT_LB = 70
USPS_OUNCE_STEPS = ("0.25", "0.5", "0.75", "0.9999")

COLUMNS = {
    "services": [
        "carrier", "service_id", "display_name", "tier", "residential_allowed",
        "business_allowed", "rate_cards", "rate_card_notes", "max_weight_lb", "notes",
    ],
    "base_rates": [
        "carrier", "service_id", "rate_card", "year", "effective_date", "zone",
        "max_weight_lb", "price", "source_file", "source_ref",
    ],
    "exceptions": [
        "carrier", "service_id", "rate_card", "year", "zone", "max_weight_lb",
        "kind", "note", "source_file", "source_ref",
    ],
    "surcharges": [
        "carrier", "year", "rate_card", "surcharge_id", "display_name", "condition",
        "applies_to", "address_type", "min_weight_lb", "max_weight_lb", "min_length_in",
        "min_width_in", "min_length_plus_girth_in", "min_volume_cuft", "zone_min", "zone_max",
        "amount", "unit", "exclusive_group", "effective_from", "effective_to",
        "fuel_applies", "source_file", "source_page", "source_quote", "notes",
    ],
    "fuel_surcharge": [
        "carrier", "year", "service_group", "default_percent", "basis",
        "window_start", "window_end", "is_proxy", "source_file", "source_ref",
        "source_quote", "notes",
    ],
    "announced_increase": [
        "carrier", "year", "product_scope", "rate_card", "percent", "source_type",
        "source_file", "source_ref", "source_quote",
    ],
    "dim_weight": [
        "carrier", "year", "effective_from", "effective_to", "divisor",
        "min_volume_cuin", "rounding", "source_file", "source_ref", "source_quote",
    ],
    "zone_from_distance": [
        "carrier", "zone", "min_miles", "max_miles", "source_file", "source_quote", "notes",
    ],
}
ADDRESS_TYPES = ("any", "residential", "commercial")
CURATED_TABLES = ("surcharges", "fuel_surcharge", "announced_increase", "dim_weight")
EXCEPTION_KINDS = (
    "ak_hi_zone", "rounding_artifact", "non_monotonic", "missing_cell",
    "page_shift", "step_structure_change", "other",
)
SURCHARGE_CONDITIONS = (
    "residential", "delivery_area", "delivery_area_extended",
    "additional_handling_weight", "additional_handling_length",
    "additional_handling_width", "additional_handling_packaging", "oversize",
    "nonstandard_length", "nonstandard_volume", "peak_residential",
    "peak_additional_handling", "peak_oversize", "peak_base",
    "transportation_time_limited",
)


def money(value: Decimal | str | float) -> str:
    """Format a price with exactly two decimals, rounding half up."""
    return str(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def weight_label(value: Decimal | str | float) -> str:
    """Normalize a weight step label: integers as '12', ounce steps as in USPS_OUNCE_STEPS."""
    dec = Decimal(str(value)).normalize()
    if dec == dec.to_integral_value():
        return str(int(dec))
    text = format(dec, "f")
    return text if text in USPS_OUNCE_STEPS else text


@lru_cache(maxsize=1)
def manifest_rows() -> list[dict]:
    with MANIFEST.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def manifest_paths() -> set[str]:
    return {row["path"] for row in manifest_rows()}


def raw_path(relative: str) -> Path:
    """Return the on-disk path of a raw file, insisting it is in the manifest and present."""
    if relative not in manifest_paths():
        raise FileNotFoundError(f"{relative} is not listed in data/raw/manifest.csv")
    path = RAW_DIR / relative
    if not path.is_file():
        raise FileNotFoundError(f"{relative} is missing; run `python3 data/fetch_raw.py`")
    return path


def processed_path(carrier: str, table: str) -> Path:
    if carrier == "shared":
        return PROCESSED_DIR / f"{table}.csv"
    return PROCESSED_DIR / carrier / f"{table}.csv"


def write_table(carrier: str, table: str, rows: Iterable[dict]) -> Path:
    """Write rows with the schema's column order. Missing keys become empty strings;
    unknown keys are an error so typos cannot slip through."""
    columns = COLUMNS[table]
    path = processed_path(carrier, table)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    for row in rows:
        unknown = set(row) - set(columns)
        if unknown:
            raise ValueError(f"{table}: unknown columns {sorted(unknown)}")
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})
    return path


def read_table(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


_WHITESPACE = re.compile(r"\s+")


def normalize_text(text: str) -> str:
    """Collapse whitespace and unify quotes/dashes so quotes can be matched across PDF line breaks."""
    text = text.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    text = text.replace("–", "-").replace("—", "-").replace(" ", " ")
    return _WHITESPACE.sub(" ", text).strip().lower()


@lru_cache(maxsize=4)
def _open_pdf(path: Path):
    import pdfplumber

    return pdfplumber.open(path)


@lru_cache(maxsize=2048)
def pdf_page_text(path: Path, page: int) -> str:
    """Text of a 1-based page index, via pdfplumber. Open documents and extracted
    pages are cached, so citing many rows from one guide stays fast."""
    pdf = _open_pdf(path)
    if page < 1 or page > len(pdf.pages):
        raise IndexError(f"{path.name} has {len(pdf.pages)} pages, asked for {page}")
    return pdf.pages[page - 1].extract_text() or ""


def html_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
    return html.unescape(re.sub(r"<[^>]+>", " ", raw))


def xlsx_text(path: Path, sheet: str | None = None) -> str:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [book[sheet]] if sheet and sheet in book.sheetnames else book.worksheets
    parts: list[str] = []
    for ws in sheets:
        for row in ws.iter_rows(values_only=True):
            parts.extend(str(cell) for cell in row if cell is not None)
    return " ".join(parts)


def source_text(source_file: str, source_ref: str | int | None) -> str:
    """Text of the cited location: PDF page, XLSX sheet, or whole HTML/CSV file."""
    path = raw_path(source_file)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return pdf_page_text(path, int(source_ref))
    if suffix == ".xlsx":
        return xlsx_text(path, str(source_ref) if source_ref else None)
    if suffix in {".htm", ".html"}:
        return html_text(path)
    return path.read_text(encoding="utf-8", errors="replace")
